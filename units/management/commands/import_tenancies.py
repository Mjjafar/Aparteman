import jdatetime
import openpyxl
from django.core.management.base import BaseCommand, CommandError

from units.models import TenancyHistory, Unit


def month_start(year, month):
    return jdatetime.date(int(year), int(month), 1).togregorian()


def next_month(year, month):
    if month == 12:
        return year + 1, 1
    return year, month + 1


def import_tenancies_from_sheet1(path, dry_run=False):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise CommandError(f"فایل پیدا نشد: {path}")
    if "Sheet1" not in wb.sheetnames:
        raise CommandError("شیت Sheet1 در فایل پیدا نشد.")
    ws = wb["Sheet1"]

    units = {u.number: u for u in Unit.objects.all()}
    months_seen = {}
    errors = []

    for lineno, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if lineno == 1:
            continue
        cells = (list(row) + [None] * 9)[:9]
        kind_code, year, month, unit_no, name = cells[0], cells[1], cells[2], cells[3], cells[4]
        if kind_code is None and year is None and unit_no is None:
            continue
        if str(kind_code).strip() != "INC-001":
            continue
        try:
            unit_no, year, month = int(unit_no), int(year), int(month)
        except (TypeError, ValueError):
            continue
        if unit_no not in units or not 1 <= month <= 12:
            continue
        name = " ".join(str(name).split()) if name else ""
        if not name:
            continue
        try:
            month_start(year, month)
        except ValueError:
            errors.append(f"سطر {lineno}: سال/ماه معتبر نیست.")
            continue
        months_seen.setdefault((unit_no, name), set()).add((year, month))

    created = 0
    for (unit_no, name), months in months_seen.items():
        ordered = sorted(months)
        start = ordered[0]
        prev = ordered[0]
        for cur in ordered[1:] + [None]:
            if cur is not None and next_month(*prev) == cur:
                prev = cur
                continue
            first = month_start(*start)
            last = month_start(*next_month(*prev))
            if not dry_run:
                TenancyHistory.objects.update_or_create(
                    unit=units[unit_no], name=name, start_date=first,
                    defaults={"phone": "", "end_date": last},
                )
            created += 1
            if cur is not None:
                start = cur
                prev = cur

    return {"created": created, "errors": errors}


class Command(BaseCommand):
    help = "بازسازی سوابق سکونت از روی نام‌های Sheet1"

    def add_arguments(self, parser):
        parser.add_argument("path", help="مسیر فایل xlsx")
        parser.add_argument("--dry-run", action="store_true", help="فقط بررسی بدون ذخیره")

    def handle(self, *args, **options):
        result = import_tenancies_from_sheet1(options["path"], dry_run=options["dry_run"])
        mode = "بررسی آزمایشی" if options["dry_run"] else "واقعی"
        self.stdout.write(f"[{mode}] بازه سکونت: {result['created']}")
        for err in result["errors"]:
            self.stderr.write(f"خطا: {err}")
