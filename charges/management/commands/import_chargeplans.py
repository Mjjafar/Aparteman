import openpyxl
from django.core.management.base import BaseCommand, CommandError

from charges.models import ChargePlan
from units.models import Unit


def import_chargeplans_from_sheet2(path, dry_run=False):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise CommandError(f"فایل پیدا نشد: {path}")
    if "Sheet2" not in wb.sheetnames:
        raise CommandError("شیت Sheet2 در فایل پیدا نشد.")
    ws = wb["Sheet2"]

    units = {u.number: u for u in Unit.objects.all()}
    created, updated, skipped, errors = 0, 0, 0, []

    for lineno, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if lineno == 1:
            continue
        unit_no, year, m1, m2, _name, amount = (list(row) + [None] * 6)[:6]
        if unit_no is None and year is None:
            skipped += 1
            continue
        try:
            unit_no = int(unit_no)
            year = int(year)
            m1 = int(m1)
            m2 = int(m2)
        except (TypeError, ValueError):
            errors.append(f"سطر {lineno}: شماره واحد/سال/ماه عدد نیست.")
            continue
        if unit_no not in units:
            errors.append(f"سطر {lineno}: واحد {unit_no} در سامانه نیست.")
            continue
        if not (1 <= m1 <= 12 and 1 <= m2 <= 12):
            errors.append(f"سطر {lineno}: ماه باید بین ۱ تا ۱۲ باشد.")
            continue
        if m2 < m1:
            errors.append(f"سطر {lineno}: ماه پایان قبل از ماه شروع است.")
            continue
        try:
            amount = int(amount)
        except (TypeError, ValueError):
            errors.append(f"سطر {lineno}: مبلغ معتبر نیست.")
            continue
        if amount <= 0:
            errors.append(f"سطر {lineno}: مبلغ باید بیشتر از صفر باشد.")
            continue
        if dry_run:
            created += 1
            continue
        _obj, was_created = ChargePlan.objects.update_or_create(
            unit=units[unit_no], year=year, start_month=m1, end_month=m2,
            defaults={"amount": amount},
        )
        if was_created:
            created += 1
        else:
            updated += 1

    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors}


class Command(BaseCommand):
    help = "وارد کردن جدول حق شارژ از Sheet2 فایل اکسل"

    def add_arguments(self, parser):
        parser.add_argument("path", help="مسیر فایل xlsx")
        parser.add_argument("--dry-run", action="store_true", help="فقط بررسی بدون ذخیره")

    def handle(self, *args, **options):
        result = import_chargeplans_from_sheet2(options["path"], dry_run=options["dry_run"])
        mode = "بررسی آزمایشی" if options["dry_run"] else "واقعی"
        self.stdout.write(
            f"[{mode}] جدید: {result['created']}، به‌روزرسانی: {result['updated']}، "
            f"رد شده: {result['skipped']}"
        )
        for err in result["errors"]:
            self.stderr.write(f"خطا: {err}")
        if result["errors"]:
            raise CommandError(f"{len(result['errors'])} سطر خطا داشت.")
