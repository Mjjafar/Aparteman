import jdatetime
import openpyxl
from django.core.management.base import BaseCommand, CommandError

from expenses.models import Expense, ExpenseCategory


def parse_jalali_date(raw):
    if raw is None:
        return None
    s = str(raw).strip().replace(" ", "")
    if not s:
        return None
    try:
        if "/" in s:
            parts = [int(p) for p in s.split("/")]
            if len(parts) != 3:
                return None
            y, m, d = parts
        elif len(s) == 8 and s.isdigit():
            y, m, d = int(s[:4]), int(s[4:6]), int(s[6:8])
        else:
            return None
        return jdatetime.date(y, m, d).togregorian()
    except (ValueError, TypeError):
        return None


def parse_amount(raw):
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    try:
        value = int(str(raw).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None
    return value if value >= 0 else None


def import_expenses_from_sheet2(path, dry_run=False):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        raise CommandError(f"فایل پیدا نشد: {path}")
    if "Sheet2" not in wb.sheetnames:
        raise CommandError("شیت Sheet2 در فایل پیدا نشد.")
    ws = wb["Sheet2"]

    categories = {c.code: c for c in ExpenseCategory.objects.all()}
    created, skipped, errors = 0, 0, []

    for lineno, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if lineno == 1:
            continue
        cells = (list(row) + [None] * 8)[:8]
        _row_no, code, title, note, bill_id, payment_id, raw_date, amount = cells
        if lineno == 209:
            raw_date, amount = "1403/07/30", 0
            note = "از طرف آقای جعفرآبادی پرداخت شد بخاطر صاف کردن بدهی خرید یک سطل چسب اضافه آمده از ساختمان"
        if lineno == 200:
            raw_date = "1403/07/30"
        if lineno == 232 and isinstance(raw_date, str) and " و " in raw_date:
            raw_date = raw_date.split(" و ")[0].strip()
        if code is None and raw_date is None and amount is None:
            skipped += 1
            continue
        code = str(code).strip() if code else ""
        if not code:
            errors.append(f"سطر {lineno}: کد هزینه خالی است.")
            continue
        if code not in categories:
            errors.append(f"سطر {lineno}: کد «{code}» در جدول انواع هزینه نیست.")
            continue
        amount = parse_amount(amount)
        if amount is None:
            errors.append(f"سطر {lineno}: مبلغ معتبر نیست.")
            continue
        spent_at = parse_jalali_date(raw_date)
        if raw_date is not None and str(raw_date).strip() and spent_at is None:
            errors.append(f"سطر {lineno}: تاریخ «{raw_date}» معتبر نیست.")
            continue
        description = str(note).strip() if note and str(note).strip() else ""
        if not description:
            description = str(title).strip() if title else ""
        if dry_run:
            created += 1
            continue
        Expense.objects.create(
            category=categories[code],
            spent_at=spent_at,
            amount=amount,
            bill_id=str(bill_id).strip() if bill_id else "",
            payment_id=str(payment_id).strip() if payment_id else "",
            description=description or categories[code].title,
        )
        created += 1

    return {"created": created, "skipped": skipped, "errors": errors}


class Command(BaseCommand):
    help = "وارد کردن هزینه‌ها از Sheet2 فایل اکسل"

    def add_arguments(self, parser):
        parser.add_argument("path", help="مسیر فایل xlsx")
        parser.add_argument("--dry-run", action="store_true", help="فقط بررسی بدون ذخیره")

    def handle(self, *args, **options):
        result = import_expenses_from_sheet2(options["path"], dry_run=options["dry_run"])
        mode = "بررسی آزمایشی" if options["dry_run"] else "واقعی"
        self.stdout.write(
            f"[{mode}] جدید: {result['created']}، رد شده: {result['skipped']}"
        )
        for err in result["errors"]:
            self.stderr.write(f"خطا: {err}")
        if result["errors"]:
            raise CommandError(f"{len(result['errors'])} سطر خطا داشت.")
